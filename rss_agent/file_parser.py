"""
Program to parse files based on the file type
"""
from docx import Document
from pypdf import PdfReader


def parse_file(file_path: str) -> str:
    file_type = file_path.strip().split(".")[-1]
    match file_type:
        case "txt":
            return txt_file_parser(file_path)
        case "pdf":
            return pdf_file_parser(file_path)
        case "docx":
            return docx_file_parser(file_path)
        case _:
            raise ValueError(f"Invalid file type: {file_path}, supported file types are txt, pdf, docx")

def docx_file_parser(file_path: str) -> str:
    doc = Document(file_path)
    content = ""
    for para in doc.paragraphs:
        content += para.text

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                content += cell.text
    return content

def txt_file_parser(file_path: str) -> str:
    content = ""
    with open(file_path, "r") as f:
        content = f.read()
    return content

def pdf_file_parser(file_path: str) -> str:
    pdf_doc = PdfReader(file_path)
    content = ""
    for page in pdf_doc.pages:
        content += page.extract_text()
    return content