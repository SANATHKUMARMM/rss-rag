import docx
from langchain_core.documents import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph

from rss_agent.embedding import search_vector_db


def write_text_file(file_path: str, content: str):
    """
    Write provided content to a text file in the provided file path
    :param file_path: file path to write
    :param content: content to write
    :return:
    """
    with open(file_path, "w") as f:
        f.write(content)


def write_docx_file(file_path: str, heading: str, content: str):
    """
    Write provided content to a docx file in the provided file path
    :param heading: heading of the docx file
    :param file_path: file path to write
    :param content: content to write
    :return:
    """
    doc = docx.Document(file_path)
    doc.add_heading(heading)
    doc.add_paragraph(content)
    doc.save(file_path)


def write_pdf_file(file_path: str, heading: str, content: str):
    """
    Write provided content to a pdf file in the provided file path
    :param file_path: file path to write
    :param heading: heading of the pdf file
    :param content: content to write
    :return:
    """
    pdf_doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
    )
    style_sheet = getSampleStyleSheet()
    story = [Paragraph(heading, style_sheet["Heading1"]), Paragraph(content, style_sheet["Normal"])]
    pdf_doc.build(story)


def fetch_chunk_vector_db(query: str) -> list[Document]:
    """
    Fetch the chunk from vector database using similarity search
    :param query: query to search vector database
    :return: list of Documents
    """
    return search_vector_db(query)


tools = [write_text_file, write_docx_file, write_pdf_file, fetch_chunk_vector_db]
